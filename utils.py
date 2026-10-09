from datetime import date, datetime
from typing import Optional

import numpy as np
import pandas as pd
from dateutil.relativedelta import relativedelta
from openpyxl import load_workbook
from pyxirr import xirr


PERIOD_MONTHS = {"annually": 12, "semiannually": 6, "quarterly": 3, "monthly": 1}


def end_of_month(date_input: date, no_of_months=0):
    new_date = date(date_input.year, date_input.month, 1)
    return new_date + relativedelta(months=no_of_months + 1) - relativedelta(days=1)


def make_df_with_interest_rows(data_list, interest=True, blank_rows=0):
    df = pd.DataFrame({"Date": data_list, "Amount*": "interest" if interest else "bullet"})

    if blank_rows > 0:
        blank_df = pd.DataFrame({"Date": [None] * blank_rows, "Amount*": [None] * blank_rows})
        df = pd.concat([df, blank_df], ignore_index=True)

    return df


def get_period(interest_repayment_period: str) -> Optional[int]:
    return PERIOD_MONTHS.get(interest_repayment_period)


def create_interest_payment_table(funding_date: date, maturity_date: date, payback_period: str) -> pd.DataFrame:
    interest_date = funding_date
    interests = []

    if payback_period != "lump-sum":
        no_of_months = get_period(payback_period)
        while end_of_month(interest_date, no_of_months) < maturity_date:
            interest_date = end_of_month(interest_date, no_of_months)
            interests.append(interest_date)

    interests.append(maturity_date)

    return make_df_with_interest_rows(interests, blank_rows=max(0, 12 - len(interests)))


def create_principal_payment_table(maturity_date: date):
    return make_df_with_interest_rows([maturity_date], interest=False, blank_rows=11)


def _schedule_lookup(schedule_df: pd.DataFrame) -> dict:
    """Map month-end date -> amount for the filled rows of an interest/principal schedule.

    Dates are snapped to their month end so they line up with the monthly grid
    (e.g. a maturity on the 15th is paid in that month). First entry per month wins.
    """
    lookup = {}
    for d, amount in zip(schedule_df["Date"], schedule_df["Amount*"]):
        if d is None or amount is None or pd.isna(d) or pd.isna(amount) or amount == "":
            continue
        if isinstance(d, (pd.Timestamp, datetime)):
            d = d.date()
        elif isinstance(d, str):
            d = pd.Timestamp(d).date()
        lookup.setdefault(end_of_month(d), amount)
    return lookup


def get_repayment_table(principal, interest, interest_df, principal_df, maturity_date, funding_date):
    interest_lookup = _schedule_lookup(interest_df)
    principal_lookup = _schedule_lookup(principal_df)

    first_month = end_of_month(funding_date)
    last_month = end_of_month(maturity_date)

    dates = []
    repayment_interest = []
    principals_repayment = []
    interests_due = []
    interests_paid = []
    interests_left = []
    principals_paid = []
    principals_left = []

    principal_left = principal
    interest_left = 0
    current_month = first_month

    while current_month <= last_month:
        if current_month == first_month:
            repayment_interest.append(None)
            principals_repayment.append(None)
            interests_due.append(None)
            interests_paid.append(None)
            principals_paid.append(None)
        else:
            # first accrual month counts one extra day (funding day)
            days = current_month.day + 1 if len(dates) == 1 else current_month.day
            interest_due = principal_left * interest * days / 360
            interests_due.append(interest_due)

            interest_repayment = interest_lookup.get(current_month)
            if interest_repayment is None:
                repayment_interest.append("n/a")
                interest_paid = 0
                interest_left += interest_due
            else:
                repayment_interest.append(interest_repayment)
                if interest_repayment == "interest":
                    interest_paid = interest_due + interest_left
                else:
                    interest_paid = min(interest_due + interest_left, float(interest_repayment))
                interest_left = interest_due + interest_left - interest_paid
            interests_paid.append(interest_paid)

            principal_repayment = principal_lookup.get(current_month)
            if principal_repayment is None:
                principals_repayment.append("n/a")
                principal_paid = 0
            else:
                principals_repayment.append(principal_repayment)
                if principal_repayment == "bullet":
                    principal_paid = principal_left
                else:
                    principal_paid = min(float(principal_repayment), principal_left)
            principals_paid.append(principal_paid)
            principal_left -= principal_paid

        interests_left.append(interest_left)
        principals_left.append(principal_left)
        dates.append(current_month)
        current_month = end_of_month(current_month, 1)

    gc_df = pd.DataFrame({
        "Month (calendar)": dates,
        "Repayment Interest": repayment_interest,
        "Repayment PV": principals_repayment,
        "Interest Due (net)": pd.array(interests_due, dtype="float64"),
        "Interest Paid": pd.array(interests_paid, dtype="float64"),
        "Interest Left": pd.array(interests_left, dtype="float64"),
        "Principal Paid": pd.array(principals_paid, dtype="float64"),
        "Principal Left": pd.array(principals_left, dtype="float64"),
    })
    gc_df["GC"] = gc_df["Interest Paid"] + gc_df["Principal Paid"]

    return gc_df


def _day_counts(dates: pd.Series) -> np.ndarray:
    """Days elapsed since the previous row (first element is 0)."""
    return pd.to_datetime(dates).diff().dt.days.fillna(0).to_numpy()


def _amortize(dates: pd.Series, flows: np.ndarray, rate: float, opening: float, zero_below: float = 0.0):
    """Roll an opening balance forward at `rate` p.a., amortizing it by `flows[1:]`.

    Returns (principal BoP, interest, amortization, principal EoP); first row is the opening position.
    A previous EoP with |value| < zero_below is treated as fully repaid.
    """
    n = len(flows)
    growth = (1 + rate) ** (_day_counts(dates) / 365) - 1

    bop = np.full(n, np.nan)
    interest = np.full(n, np.nan)
    amortization = np.full(n, np.nan)
    eop = np.full(n, np.nan)
    eop[0] = opening

    for i in range(1, n):
        prev = eop[i - 1]
        bop[i] = 0.0 if abs(prev) < zero_below else prev
        interest[i] = bop[i] * growth[i]
        amortization[i] = flows[i] - interest[i]
        eop[i] = bop[i] - amortization[i]

    return bop, interest, amortization, eop


def create_cf_table(gc, wht, agent_fee, total_investment, Annual_Interest_Rate):
    cf = gc[["Month (calendar)", "Interest Paid", "Principal Paid"]].copy()
    cf["WHT on Loan Interest"] = -wht * cf["Interest Paid"]
    cf["GC after costs after cap after cuts"] = cf["Interest Paid"] + cf["Principal Paid"] + cf["WHT on Loan Interest"]
    cf["Agent Fee"] = agent_fee * cf["GC after costs after cap after cuts"]
    cf["Net Cash flow before tax"] = cf["GC after costs after cap after cuts"] - cf["Agent Fee"]
    cf.loc[cf.index[0], "Net Cash flow before tax"] = -total_investment

    net_cf = cf["Net Cash flow before tax"]
    cf["Cumulative Net CF"] = net_cf.cumsum()
    multiple = (net_cf.cumsum() - net_cf.iloc[0]) / total_investment
    multiple.iloc[0] = np.nan
    cf["Net multiple for Success fee"] = multiple

    # --- loan waterfall: net CF split into interest / principal / variable interest ---
    n = len(cf)
    net = net_cf.to_numpy(dtype=float)
    daily_rate = (1 + Annual_Interest_Rate) ** (1 / 365) - 1
    growth = (1 + daily_rate) ** _day_counts(cf["Month (calendar)"]) - 1

    interest_due = np.full(n, np.nan)
    interest_paid = np.full(n, np.nan)
    interest_left = np.full(n, np.nan)
    principal_paid = np.full(n, np.nan)
    principal_left = np.full(n, np.nan)
    variable_interest = np.full(n, np.nan)
    interest_left[0] = 0
    principal_left[0] = total_investment

    for i in range(1, n):
        interest_due[i] = principal_left[i - 1] * growth[i]
        interest_paid[i] = max(min(net[i], interest_due[i] + interest_left[i - 1]), 0)
        interest_left[i] = interest_left[i - 1] + interest_due[i] - interest_paid[i]
        principal_paid[i] = min(net[i] - interest_paid[i], principal_left[i - 1])
        principal_left[i] = principal_left[i - 1] - principal_paid[i]
        variable_interest[i] = net[i] - principal_paid[i] - interest_paid[i]

    cf["Interest Due (net)"] = interest_due
    cf["Interest Paid - net"] = interest_paid
    cf["Interest Left"] = interest_left
    cf["Principal Paid - net"] = principal_paid
    cf["Principal Left"] = principal_left
    cf["Variable Interest"] = variable_interest

    return cf


def create_npv_table(cf, total_investment):
    npv_local_spv = cf[["Month (calendar)", "GC after costs after cap after cuts"]].rename(
        columns={"GC after costs after cap after cuts": "TECH for gross IRR"})
    npv_local_spv.loc[npv_local_spv.index[0], "TECH for gross IRR"] = -total_investment

    rate = xirr(npv_local_spv["Month (calendar)"], npv_local_spv["TECH for gross IRR"])

    bop, interest, amortization, eop = _amortize(
        npv_local_spv["Month (calendar)"], npv_local_spv["TECH for gross IRR"].to_numpy(dtype=float),
        rate, total_investment, zero_below=1)

    npv_local_spv["Principal BoP"] = bop
    npv_local_spv["Interest"] = interest
    npv_local_spv["Amortization"] = amortization
    npv_local_spv["Principal EoP/NPV"] = eop

    return npv_local_spv


def _running_irr(dates, flows) -> list:
    """XIRR of every prefix of the cash flows (-1 where it is undefined), rounded to 4 dp."""
    dates = list(dates)
    flows = np.asarray(flows, dtype=float)
    out = []
    has_positive = False
    for i in range(len(flows)):
        has_positive = has_positive or flows[i] > 0
        irr = None
        if has_positive:
            try:
                irr = xirr(dates[:i + 1], flows[:i + 1])
            except Exception:
                irr = None
        out.append(round(irr, 4) if irr is not None else -1)
    return out


def update_cf(cf, npv_local, total_investment, tax_rate, sicav_costs_monthly):
    cf = cf.copy()
    first = cf.index[0]

    cf["Tax base CIT"] = npv_local["Interest"] - cf["Agent Fee"] - cf["Interest Due (net)"]
    cf["Income tax (CIT)"] = cf["Tax base CIT"] * tax_rate
    cf["Net CF SICAV incoming"] = cf["Net Cash flow before tax"] - cf["Income tax (CIT)"]
    cf.loc[first, "Net CF SICAV incoming"] = -total_investment

    sicav_npv = cf[["Month (calendar)", "Net CF SICAV incoming"]].copy()
    rate_sicav = xirr(sicav_npv["Month (calendar)"], sicav_npv["Net CF SICAV incoming"])

    bop, interest, amortization, eop = _amortize(
        sicav_npv["Month (calendar)"], sicav_npv["Net CF SICAV incoming"].to_numpy(dtype=float),
        rate_sicav, total_investment)

    sicav_npv["Principal BoP"] = bop
    sicav_npv["Interest"] = interest
    sicav_npv["Amortization"] = amortization
    sicav_npv["Principal EoP/NPV"] = eop

    cf["SICAV Costs"] = sicav_costs_monthly * sicav_npv["Principal BoP"]
    cf["Net CF after all taxes"] = cf["Net CF SICAV incoming"] - cf["SICAV Costs"]
    cf.loc[first, "Net CF after all taxes"] = -total_investment

    cf["IRR"] = _running_irr(cf["Month (calendar)"], cf["Net CF after all taxes"])
    cf["Cumulative Net CF after all taxes"] = cf["Net CF after all taxes"].cumsum()
    cf["Net break even point // TEX"] = (cf["IRR"] > 0).astype(int)

    return cf, sicav_npv


def create_all_tables(gc, wht, agent_fee, total_investment, Annual_Interest_Rate, tax_rate, sicav_costs_monthly):
    cf = create_cf_table(gc=gc, wht=wht, agent_fee=agent_fee, total_investment=total_investment, Annual_Interest_Rate=Annual_Interest_Rate)
    npv_table = create_npv_table(cf, total_investment)
    cf, sicav_npv = update_cf(cf=cf, npv_local=npv_table, total_investment=total_investment, tax_rate=tax_rate, sicav_costs_monthly=sicav_costs_monthly)

    return cf, npv_table, sicav_npv


def save_model(input_file_name, **kwargs):
    wb = load_workbook(filename=f"{input_file_name}.xlsx")

    ws = wb["GC"]

    ws["D9"] = kwargs["funding_amount"]
    ws["D10"] = kwargs["funding_date"]
    ws["D11"] = kwargs["maturity_date"]
    ws["D12"] = kwargs["interest"]
    ws["D13"] = kwargs["repayment_period"]

    for i in range(12):
        ws[f"G{i + 7}"] = kwargs["interest_df"]["Date"].iloc[i]
        ws[f"H{i + 7}"] = kwargs["interest_df"]["Amount*"].iloc[i]

        ws[f"J{i + 7}"] = kwargs["principal_df"]["Date"].iloc[i]
        ws[f"K{i + 7}"] = kwargs["principal_df"]["Amount*"].iloc[i]

    ws = wb["Summary&CF_SICAV"]

    ws["D3"] = kwargs["project_name"]

    ws["I86"] = kwargs["engagement_fee"]
    ws["I86"] = kwargs["arrangement_fee"]

    ws["E24"] = kwargs["wht_interest"]
    ws["E25"] = kwargs["agent_fee"]
    ws["E48"] = kwargs["SICAV_costs"]

    return wb
